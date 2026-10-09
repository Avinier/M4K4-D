"""RP-01 fullproofmath §11.3 busy minute, re-run for the D-049 head.

Reads the BC-60 and MV-60 5 ms traces of the RP-01 workbook
(layout03-paper-model.xlsx). The nominal torque is tau = J * alpha (A0 puts
the CoM on the axes; zero bias), so every untouched segment is reused as
authored. The D-046 storyboard delta (feetech-motion-storyboard.md) stretches
the best-case laugh reversals and both startle recoils: those strokes are
rest to rest on the 5 ms grid, so their law (MJ5 or MS7) is fitted to the
workbook samples and re-sampled over the new duration; the neutral hold that
follows absorbs the extra time, so each minute stays 60 s.

Validation: the unmodified traces with the workbook's inertias reproduce its
Summary RMS/peaks. Output-side rigid-body torque only; the STS3045M reflected
motor inertia (281:1, U) and friction are excluded, as in the workbook.
Needs openpyxl. Writes generated/busy-minute-d049.json.
"""
import json
from pathlib import Path
import numpy as np
import openpyxl

HERE = Path(__file__).parent
BOOK = HERE.parents[3] / '02-prototypes' / 'RP-01-head' / 'layout03-paper-model.xlsx'
COLS = {'pitch': (2, 3, 4), 'roll': (7, 8, 9), 'yaw': (12, 13, 14)}   # angle, speed, accel
DT = 0.005


def mj5(s): return 10*s**3 - 15*s**4 + 6*s**5, 30*s**2 - 60*s**3 + 30*s**4, 60*s - 180*s**2 + 120*s**3
def ms7(s): return 35*s**4 - 84*s**5 + 70*s**6 - 20*s**7, 140*s**3 - 420*s**4 + 420*s**5 - 140*s**6, 420*s**2 - 1680*s**3 + 2100*s**4 - 840*s**5
LAWS = {'MJ5': mj5, 'MS7': ms7}

wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)


def trace(sheet):
    rows = [r for r in wb[sheet].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    return [dict(t=r[0], seg=r[1], **{f'{ax}_{k}': r[c] for ax, cs in COLS.items() for k, c in zip(('a', 'v', 'acc'), cs)}) for r in rows]


def runs(rows):
    out, start = [], 0
    for i in range(1, len(rows) + 1):
        if i == len(rows) or rows[i]['seg'] != rows[start]['seg']:
            out.append([rows[start]['seg'], start, i]); start = i
    return out


def stroke(rows, i0, i1, ms):
    """Re-sample one rest-to-rest stroke (rows i0..i1-1) over a new duration."""
    n_old = i1 - i0; T_old = n_old * DT
    new = []
    n_new = int(np.ceil(ms / DT - 1e-9))   # round up: never faster than the 70% cap
    laws = {}
    for ax in COLS:
        a0, a1 = rows[i0][f'{ax}_a'], rows[i1][f'{ax}_a']
        if abs(a1 - a0) < 1e-9: laws[ax] = (a0, a1, None); continue
        s = np.arange(n_old) / n_old
        old = np.array([rows[i0 + k][f'{ax}_a'] for k in range(n_old)])
        fits = {k: np.max(np.abs(a0 + (a1 - a0) * f(s)[0] - old)) for k, f in LAWS.items()}
        law = min(fits, key=fits.get)
        assert fits[law] < 1e-6, (rows[i0]['seg'], ax, fits)
        laws[ax] = (a0, a1, law)
    T = n_new * DT
    for k in range(n_new):
        s = k / n_new; r = dict(t=None, seg=rows[i0]['seg'])
        for ax, (a0, a1, law) in laws.items():
            if law is None:
                r.update({f'{ax}_a': a0, f'{ax}_v': 0., f'{ax}_acc': 0.}); continue
            p, v, a = LAWS[law](s)
            r.update({f'{ax}_a': a0 + (a1 - a0) * p, f'{ax}_v': (a1 - a0) * v / T,
                      f'{ax}_acc': np.radians(a1 - a0) * a / T**2})
        new.append(r)
    return new, {ax: l[2] for ax, l in laws.items() if l[2]}, round(T_old, 3), round(T, 3)


def retime(rows, plan):
    """plan: {segment label: new duration s}; each stretch is absorbed by the
    next 'Neutral hold' run (each re-timed phrase is followed by one)."""
    out, log, debt = [], [], 0
    for label, i0, i1 in runs(rows):
        block = rows[i0:i1]
        if label in plan:
            block, laws, t_old, t_new = stroke(rows, i0, i1, plan[label])
            debt += len(block) - (i1 - i0)
            log.append(dict(segment=label, laws=laws, old_s=t_old, new_s=t_new))
        elif debt and label == 'Neutral hold':
            assert len(block) > debt, f'{label} cannot absorb {debt * DT:.3f} s'
            block = block[debt:]; debt = 0
        out.extend(block)
    assert debt == 0, 'unabsorbed stretch'
    for k, r in enumerate(out): r['t'] = round(k * DT, 6)
    return out, log


def results(rows, J):
    res = {}
    for ax in COLS:
        tau = np.array([J[ax] * r[f'{ax}_acc'] for r in rows])
        k = int(np.argmax(np.abs(tau)))
        res[ax] = dict(rms_Nm=float(np.sqrt(np.mean(tau ** 2))), peak_Nm=float(abs(tau[k])),
                       speed_at_peak_deg_s=float(rows[k][f'{ax}_v']), time_s=rows[k]['t'], segment=rows[k]['seg'],
                       peak_speed_deg_s=float(max(abs(r[f'{ax}_v']) for r in rows)))
    return res


assumptions = {r[0]: r[1] for r in wb['Assumptions'].iter_rows(min_row=6, values_only=True) if r[0]}
summary = {r[0]: dict(rms=r[1], peak=r[3]) for r in wb['Summary'].iter_rows(min_row=6, max_row=11, values_only=True)}
J_old = dict(pitch=assumptions['Pitch inertia'], roll=assumptions['Roll inertia'], yaw=assumptions['Yaw inertia'])
tree = json.loads((HERE / 'mass-placement.json').read_text())['working']
J_new = {ax: tree[ax]['estimated_inertia_kg_m2'] for ax in COLS}
PLAN = {'BC-60': {**{f'HM-08 Laugh {p} {n}': d for p in 'AB' for n, d in ((2, .143), (3, .172), (4, .158))},
                  'HM-15 Startle recoil': .258},
        'MV-60': {'HM-15 Startle recoil': .358}}
report = dict(method=__doc__.strip().splitlines()[0], workbook=str(BOOK.relative_to(HERE.parents[3])),
              inertia_workbook_kg_m2=J_old, inertia_d049_kg_m2=J_new, cases={})
ok = True
for case in ('BC-60', 'MV-60'):
    rows = trace(case)
    base = results(rows, J_old)
    val = {}
    for ax in COLS:
        ref = summary[f'{case} {ax.capitalize()}']
        val[ax] = dict(workbook_rms=ref['rms'], rebuilt_rms=base[ax]['rms_Nm'], workbook_peak=ref['peak'], rebuilt_peak=base[ax]['peak_Nm'])
        ok &= abs(ref['rms'] - base[ax]['rms_Nm']) < 1e-9 and abs(ref['peak'] - base[ax]['peak_Nm']) < 1e-9
    new, log = retime(rows, PLAN[case])
    assert abs(len(new) - len(rows)) == 0
    report['cases'][case] = dict(validation_against_workbook=val, retimed=log,
                                 d046_trajectory_workbook_inertia=results(new, J_old), d049=results(new, J_new))
# Moving-curve margins for the STS3045M axes (feetech-actuator-screen.md
# method): straight-line proxy at the 5.23 V worst head terminal x 0.86,
# demand (J_out + J_eq) * alpha with the screen's reflected-inertia range,
# plus the §11.3 severe 0.020 N*m bias corner added to |tau|.
V, V0 = 5.23, 6.0
T_STALL, RPM0, FACTOR = 0.588 * V / V0, 75. * V / V0, .86
J_EQ = (0.000869, 0.001974)
BIAS = (0., .020)


def margins(rows, J):
    out = {}
    for ax in ('pitch', 'roll'):
        acc = np.array([r[f'{ax}_acc'] for r in rows]); rpm = np.abs([r[f'{ax}_v'] for r in rows]) / 6.
        cap = FACTOR * T_STALL * (1 - rpm / RPM0)
        for je in J_EQ:
            for bias in BIAS:
                need = np.abs((J[ax] + je) * acc) + bias
                k = int(np.argmin(cap / need))
                out[f'{ax}_Jeq{je}_bias{bias}'] = dict(min_ratio=float(cap[k] / need[k]), demand_Nm=float(need[k]),
                                                       capability_Nm=float(cap[k]), rpm=float(rpm[k]), segment=rows[k]['seg'])
    return out


for case in report['cases']:
    new, _ = retime(trace(case), PLAN[case])
    report['cases'][case]['sts3045m_margins'] = margins(new, J_new)
report['capability_proxy'] = dict(terminal_V=V, stall_Nm=T_STALL, no_load_rpm=RPM0, graph_factor=FACTOR,
                                  basis='E: 6 V vendor endpoints scaled to 5.23 V, straight line, legacy XC330 0.86 factor; not a Feetech curve')
report['validation_passed'] = bool(ok)
report['limitations'] = [
    'Nominal A0 case: zero CoM residual and zero external bias. The §11.3 sensitivity matrix (C2 mass, CoM offsets, +/-0.020 N*m bias) is not re-run.',
    'Output-side rigid-body torque only; add the reflected motor inertia (J_motor x 281^2, U) and friction before comparing with a measured Feetech curve.',
    'Inertias are the E mass tree (solid PLA, uniform-density servo envelopes).']
(HERE / 'generated' / 'busy-minute-d049.json').write_text(json.dumps(report, indent=2) + '\n')
for case, r in report['cases'].items():
    for ax in COLS:
        v, a, n = r['validation_against_workbook'][ax], r['d046_trajectory_workbook_inertia'][ax], r['d049'][ax]
        print(f"{case} {ax:5s} workbook {v['workbook_rms']:.6f}/{v['workbook_peak']:.6f} rebuilt {v['rebuilt_rms']:.6f}/{v['rebuilt_peak']:.6f} | "
              f"D-046 traj {a['rms_Nm']:.6f}/{a['peak_Nm']:.6f} | D-049 {n['rms_Nm']:.6f}/{n['peak_Nm']:.6f} at {n['speed_at_peak_deg_s']:+.1f} deg/s {n['segment']}")
    print(' retimed', r['retimed'])
for case, r in report['cases'].items():
    for k, v in r['sts3045m_margins'].items():
        print(f"{case} {k:28s} min ratio {v['min_ratio']:.2f} demand {v['demand_Nm']:.4f} cap {v['capability_Nm']:.4f} at {v['rpm']:.1f} rpm {v['segment']}")
print('validation', 'PASS' if ok else 'FAIL')
raise SystemExit(0 if ok else 1)
